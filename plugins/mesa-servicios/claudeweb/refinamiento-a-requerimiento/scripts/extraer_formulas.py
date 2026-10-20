#!/usr/bin/env python3
"""Inventario de fórmulas de un Excel entregado por el solicitante.

Convierte "ajusten la fórmula, va el Excel" en una lista de candidatas con
ubicación, contexto y valor, para que quien levanta el requerimiento
identifique la fórmula y la confirme con el cliente con una pregunta cerrada.

Uso:
    python extraer_formulas.py archivo.xlsx [--buscar texto ...] [--hoja Nombre] [--max 300]

Salida (Markdown):
  1. Resumen por hoja: filas, columnas, cantidad de fórmulas y de fórmulas distintas.
  2. Fórmulas agrupadas por patrón (misma fórmula arrastrada en varias filas
     cuenta como una), con: hoja, rango, encabezado de columna, etiqueta de fila,
     fórmula de ejemplo, valor calculado de ejemplo y celdas de las que depende.
  3. Con --buscar, solo los patrones cuyo encabezado, etiqueta o fórmula
     contienen alguno de los textos.
"""
import argparse
import re
import sys
from collections import OrderedDict

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REF_RE = re.compile(r"(?:'([^']+)'!|(\w[\w.]*)!)?\$?\b([A-Z]{1,3})\$?(\d+)\b")


def patron(formula, fila):
    """Normaliza referencias relativas a la fila para agrupar fórmulas arrastradas."""
    def rel(m):
        hoja, col, f = m.group(1) or m.group(2), m.group(3), int(m.group(4))
        pref = f"{hoja}!" if hoja else ""
        return f"{pref}{col}[{f - fila:+d}]"
    return REF_RE.sub(rel, formula)


def encabezado_col(ws, col, fila_formula, max_buscar=10):
    """Primer texto hacia arriba en la columna (encabezado más cercano)."""
    for r in range(fila_formula - 1, max(0, fila_formula - max_buscar - 1), -1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()
    for r in range(1, 6):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()
    return ""


def etiqueta_fila(ws, fila, col_formula):
    """Primer texto a la izquierda en la fila."""
    for c in range(col_formula - 1, 0, -1):
        v = ws.cell(row=fila, column=c).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()
    return ""


def fmt(v):
    if v is None:
        return "—"
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.4f}".rstrip("0").rstrip(".")
    return str(v)[:40]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("archivo")
    ap.add_argument("--buscar", nargs="*", default=[])
    ap.add_argument("--hoja")
    ap.add_argument("--max", type=int, default=300)
    a = ap.parse_args()

    wf = load_workbook(a.archivo, data_only=False)
    wv = load_workbook(a.archivo, data_only=True)
    buscar = [b.lower() for b in a.buscar]

    print(f"# Inventario de fórmulas — {a.archivo}\n")
    print("| Hoja | Filas × columnas | Celdas con fórmula | Fórmulas distintas |")
    print("|---|---|---|---|")

    grupos_total = []
    for ws in wf.worksheets:
        if a.hoja and ws.title != a.hoja:
            continue
        wsv = wv[ws.title]
        grupos = OrderedDict()
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    key = (c.column, patron(c.value, c.row))
                    g = grupos.setdefault(key, {"celdas": [], "formula": c.value,
                                                "fila": c.row, "col": c.column})
                    g["celdas"].append(c.row)
        print(f"| {ws.title} | {ws.max_row} × {ws.max_column} | {n} | {len(grupos)} |")
        for (col, _), g in grupos.items():
            letra = get_column_letter(col)
            filas = g["celdas"]
            rango = f"{letra}{filas[0]}" if len(filas) == 1 else f"{letra}{filas[0]}:{letra}{filas[-1]}"
            enc = encabezado_col(ws, col, g["fila"])
            et = etiqueta_fila(ws, g["fila"], col)
            val = wsv.cell(row=g["fila"], column=col).value
            deps = sorted({((m.group(1) or m.group(2)) + "!" if (m.group(1) or m.group(2)) else "")
                           + m.group(3) + m.group(4) for m in REF_RE.finditer(g["formula"])})
            texto = f"{enc} {et} {g['formula']}".lower()
            if buscar and not any(b in texto for b in buscar):
                continue
            grupos_total.append((ws.title, rango, len(filas), enc, et, g["formula"], val, deps))

    if not grupos_total:
        print("\nSin fórmulas" + (" que coincidan con la búsqueda." if buscar else "."))
        return

    print(f"\n## Fórmulas por patrón ({len(grupos_total)})\n")
    print("Una fórmula arrastrada en varias filas cuenta como un patrón. "
          "El valor es el calculado en la primera celda del rango (requiere que el Excel "
          "se haya guardado con valores; si aparece '—', recalcular).\n")
    print("| # | Hoja | Rango | Celdas | Encabezado | Etiqueta de fila | Fórmula (ejemplo) | Valor | Depende de |")
    print("|---|---|---|---|---|---|---|---|---|")
    for i, (h, r, n, e, et, f, v, d) in enumerate(grupos_total[:a.max], 1):
        f_md = f.replace("|", "\\|")
        print(f"| F-{i:02d} | {h} | {r} | {n} | {e} | {et} | `{f_md}` | {fmt(v)} | {', '.join(d[:6])}{'…' if len(d) > 6 else ''} |")
    if len(grupos_total) > a.max:
        print(f"\n… {len(grupos_total) - a.max} patrones más. Acotar con --buscar o --hoja.")


if __name__ == "__main__":
    sys.exit(main())
